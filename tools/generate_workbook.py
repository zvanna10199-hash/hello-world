"""Generates the pastel budget / habit-tracker / task-planner workbook.

Run: python3 tools/generate_workbook.py
Output: Планировщик_Бюджет_Привычки.xlsx (repo root)
"""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule, DataBarRule
from openpyxl.chart import PieChart, BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment

# ---------------------------------------------------------------- palette --
PURPLE      = "D9CCE8"
PURPLE_DK   = "B9A6D6"
SAGE        = "CDE6D0"
SAGE_DK     = "9DC3A4"
SAGE_LIGHT  = "E9F4EA"
PINK        = "F7DCE2"
PINK_DK     = "E7B8C2"
PINK_LIGHT  = "FBEDF1"
YELLOW      = "F6EBC0"
YELLOW_DK   = "EAD98A"
YELLOW_LIGHT= "FBF6E3"
BLUE        = "D3E6F0"
BLUE_DK     = "A9C8DE"
BLUE_LIGHT  = "EBF3F8"
CREAM       = "FBF7EF"
PEACH       = "F3D9C4"
GRAY_TXT    = "5B5B5B"
GRAY_LIGHT  = "EDEDED"
WHITE       = "FFFFFF"

FONT_TITLE = Font(name="Roboto", size=16, bold=True, color=GRAY_TXT)
FONT_H2    = Font(name="Roboto", size=11, bold=True, color=GRAY_TXT)
FONT_LBL   = Font(name="Roboto", size=9, bold=True, color=GRAY_TXT)
FONT_BODY  = Font(name="Roboto", size=10, color=GRAY_TXT)
FONT_BIG   = Font(name="Roboto", size=18, bold=True, color=GRAY_TXT)
FONT_NOTE  = Font(name="Roboto", size=9, italic=True, color=GRAY_TXT)

THIN = Side(style="thin", color="D9D9D9")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def fill(color):
    return PatternFill("solid", fgColor=color)

def style_title(ws, rng, text, color=PURPLE, font=None):
    ws.merge_cells(rng)
    cell = ws[rng.split(":")[0]]
    cell.value = text
    cell.font = font or FONT_TITLE
    cell.fill = fill(color)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws[rng]:
        for c in row:
            c.fill = fill(color)

def box_header(ws, rng, text, color):
    ws.merge_cells(rng)
    cell = ws[rng.split(":")[0]]
    cell.value = text
    cell.font = FONT_LBL
    cell.fill = fill(color)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws[rng]:
        for c in row:
            c.fill = fill(color)
            c.border = BORDER_ALL

def set_col_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

wb = Workbook()

# =====================================================================
# SHEET: СПРАВОЧНИК (reference lists, must exist first for validations)
# =====================================================================
ref = wb.active
ref.title = "Справочник"
ref.sheet_properties.tabColor = PURPLE_DK
set_col_widths(ref, {"A": 22, "B": 22, "C": 24, "D": 16, "E": 18, "F": 14, "G": 16, "H": 24})

style_title(ref, "A1:H1", "СПРАВОЧНИК КАТЕГОРИЙ И СПИСКОВ", PURPLE)
ref.row_dimensions[1].height = 26

headers = ["Категории доходов", "Категории расходов", "Служебные (переводы/долг/инвест.)",
           "Категории задач", "Приоритет", "Статус задачи"]
for i, h in enumerate(headers):
    c = ref.cell(row=2, column=1 + i, value=h)
    c.font = FONT_LBL
    c.fill = fill(BLUE_LIGHT)
    c.border = BORDER_ALL
    c.alignment = Alignment(horizontal="center", wrap_text=True)

income_cats = ["Зарплата", "Заказы", "Подарок"]
expense_cats = ["Аренда", "Продукты", "Обучение", "Кредит", "Субподряд", "Ипотека", "Подарок", "Аутсорс"]
service_cats = ["Взял в долг", "Вернул долг", "Дал в долг", "Вернули долг",
                "Пополнение вклада", "Снятие со вклада", "Инвестиции: пополнение", "Инвестиции: вывод"]
task_cats = ["Работа", "Личное", "Другое"]
priorities = ["Срочно", "Высокий", "Средний", "Низкий"]
statuses = ["Не начато", "В процессе", "Готово"]
accounts = ["Наличные", "Карта Сбербанк", "Карта Тинькофф", "Вклад", "Кредит 1", "Кредит 2", "Инвестиционный счёт"]
month_names = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
               "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
years_list = list(range(2023, 2033))

cols_data = [income_cats, expense_cats, service_cats, task_cats, priorities, statuses]
for ci, data in enumerate(cols_data):
    for ri, val in enumerate(data):
        cell = ref.cell(row=3 + ri, column=1 + ci, value=val)
        cell.font = FONT_BODY
        cell.border = BORDER_ALL

note_row = 3 + max(len(c) for c in cols_data) + 1
ref.merge_cells(start_row=note_row, start_column=3, end_row=note_row + 3, end_column=3)
note_cell = ref.cell(row=note_row, column=3)
note_cell.value = ("⚠ Эти категории — служебные (переводы между своими счетами, долги, инвестиции).\n"
                    "Они не считаются доходом/расходом в сводке за месяц, но участвуют в остатках по счетам.\n"
                    "Не изменяйте и не удаляйте их, если хотите корректный учёт.")
note_cell.font = FONT_NOTE
note_cell.fill = fill(YELLOW_LIGHT)
note_cell.alignment = Alignment(wrap_text=True, vertical="top")

# combined category list used by the "Категория" dropdown on Бюджет sheet
ref["J2"] = "Все категории (для выпадающего списка)"
ref["J2"].font = FONT_LBL
all_cats = income_cats + [c for c in expense_cats if c not in income_cats] + service_cats
for i, val in enumerate(all_cats):
    ref.cell(row=3 + i, column=10, value=val).font = FONT_BODY
ref.column_dimensions["J"].width = 26

ref["K2"] = "Счета / способы оплаты"
ref["K2"].font = FONT_LBL
ref["K2"].fill = fill(BLUE_LIGHT)
ref["K2"].border = BORDER_ALL
for i, val in enumerate(accounts):
    ref.cell(row=3 + i, column=11, value=val).font = FONT_BODY
ref.column_dimensions["K"].width = 20
ref.merge_cells(start_row=3 + len(accounts) + 1, start_column=11, end_row=3 + len(accounts) + 3, end_column=11)
acc_note = ref.cell(row=3 + len(accounts) + 1, column=11)
acc_note.value = "✏ Добавьте сюда свои карты, наличные, вклады, кредиты или инвестсчета — они появятся в выпадающем списке «Счёт» в Бюджете."
acc_note.font = FONT_NOTE
acc_note.fill = fill(YELLOW_LIGHT)
acc_note.alignment = Alignment(wrap_text=True, vertical="top")

ref["L2"] = "Месяцы"
ref["L2"].font = FONT_LBL
ref["L2"].fill = fill(BLUE_LIGHT)
ref["L2"].border = BORDER_ALL
for i, val in enumerate(month_names):
    ref.cell(row=3 + i, column=12, value=val).font = FONT_BODY
ref.column_dimensions["L"].width = 12

ref["M2"] = "Годы"
ref["M2"].font = FONT_LBL
ref["M2"].fill = fill(BLUE_LIGHT)
ref["M2"].border = BORDER_ALL
for i, val in enumerate(years_list):
    ref.cell(row=3 + i, column=13, value=val).font = FONT_BODY
ref.column_dimensions["M"].width = 10

credit_types = ["Кредит", "Кредитная карта"]
credit_statuses = ["Активен", "Закрыт"]
payoff_strategies = ["Снежный ком", "Лавина"]

ref["N2"] = "Тип кредита"
ref["O2"] = "Статус кредита"
ref["P2"] = "Стратегия закрытия"
for col, data in zip((14, 15, 16), (credit_types, credit_statuses, payoff_strategies)):
    c = ref.cell(row=2, column=col)
    c.font = FONT_LBL
    c.fill = fill(BLUE_LIGHT)
    c.border = BORDER_ALL
    for i, val in enumerate(data):
        ref.cell(row=3 + i, column=col, value=val).font = FONT_BODY
ref.column_dimensions["N"].width = 16
ref.column_dimensions["O"].width = 12
ref.column_dimensions["P"].width = 16

ref.sheet_view.showGridLines = False

def named_list_range(col_letter, n):
    return f"Справочник!${col_letter}$3:${col_letter}${2+n}"

RANGE_INCOME = named_list_range("A", len(income_cats))
RANGE_EXPENSE = named_list_range("B", len(expense_cats))
RANGE_SERVICE = named_list_range("C", len(service_cats))
RANGE_ALLCATS = named_list_range("J", len(all_cats))
RANGE_TASKCAT = named_list_range("D", len(task_cats))
RANGE_PRIORITY = named_list_range("E", len(priorities))
RANGE_STATUS = named_list_range("F", len(statuses))
RANGE_ACCOUNTS = named_list_range("K", len(accounts))
RANGE_MONTHS = named_list_range("L", len(month_names))
RANGE_YEARS = named_list_range("M", len(years_list))
RANGE_CREDIT_TYPES = named_list_range("N", len(credit_types))
RANGE_CREDIT_STATUSES = named_list_range("O", len(credit_statuses))
RANGE_PAYOFF_STRATEGIES = named_list_range("P", len(payoff_strategies))

def add_checkbox_validation(ws, cell_range):
    dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv

def add_list_validation(ws, cell_range, source_range):
    dv = DataValidation(type="list", formula1=f"={source_range}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv

def add_list_validation_literal(ws, cell_range, csv_values):
    dv = DataValidation(type="list", formula1=f'"{csv_values}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv

# =====================================================================
# SHEET: ТРЕКЕР ПРИВЫЧЕК
# =====================================================================
hab = wb.create_sheet("Трекер привычек")
hab.sheet_properties.tabColor = SAGE_DK
hab.sheet_view.showGridLines = False

N_DAYS = 31
FIRST_DAY_COL = 3  # column C
LAST_DAY_COL = FIRST_DAY_COL + N_DAYS - 1  # AG = 33

widths = {"A": 3, "B": 27}
for col in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
    widths[get_column_letter(col)] = 3.4
set_col_widths(hab, widths)

# Row1: title bar
style_title(hab, f"A1:{get_column_letter(LAST_DAY_COL)}1", "🌿  ТРЕКЕР ПРИВЫЧЕК", PURPLE)
hab.row_dimensions[1].height = 28

# Row2: month/year switcher
hab["A2"] = "Месяц:"
hab["A2"].font = FONT_LBL
hab["B2"] = "Июль"
hab["B2"].font = Font(name="Roboto", size=10, bold=True, color=GRAY_TXT)
hab["B2"].fill = fill(YELLOW_LIGHT)
hab["B2"].alignment = Alignment(horizontal="center")
hab["B2"].border = BORDER_ALL
add_list_validation(hab, "B2", RANGE_MONTHS)

hab["C2"] = "Год:"
hab["C2"].font = FONT_LBL
hab["D2"] = 2026
hab["D2"].font = Font(name="Roboto", size=10, bold=True, color=GRAY_TXT)
hab["D2"].fill = fill(YELLOW_LIGHT)
hab["D2"].alignment = Alignment(horizontal="center")
hab["D2"].border = BORDER_ALL
add_list_validation(hab, "D2", RANGE_YEARS)

START_CELL = f"DATE($D$2,MATCH($B$2,{RANGE_MONTHS},0),1)"

# Row4-5: month/date pink block (merged 2 rows)
hab.merge_cells("B4:B5")
month_cell = hab["B4"]
month_cell.value = f'=UPPER(TEXT({START_CELL},"mmmm"))&CHAR(10)&TEXT({START_CELL},"d mmmm yyyy")'
month_cell.font = Font(name="Roboto", size=13, bold=True, color=GRAY_TXT)
month_cell.fill = fill(PINK)
month_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for rr in (4, 5):
    hab.cell(row=rr, column=2).border = BORDER_ALL

week_blocks = [
    ("НЕДЕЛЯ 1", FIRST_DAY_COL, FIRST_DAY_COL + 6, BLUE),
    ("НЕДЕЛЯ 2", FIRST_DAY_COL + 7, FIRST_DAY_COL + 13, SAGE),
    ("НЕДЕЛЯ 3", FIRST_DAY_COL + 14, FIRST_DAY_COL + 20, PINK),
    ("НЕДЕЛЯ 4", FIRST_DAY_COL + 21, FIRST_DAY_COL + 27, YELLOW),
    ("ДОП.", FIRST_DAY_COL + 28, LAST_DAY_COL, SAGE),
]
for label, c1, c2, color in week_blocks:
    rng = f"{get_column_letter(c1)}4:{get_column_letter(c2)}4"
    box_header(hab, rng, label, color)

# Row5: weekday abbreviations, Row6: date numbers
for col in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
    idx = col - FIRST_DAY_COL
    letter = get_column_letter(col)
    wd = hab.cell(row=5, column=col)
    wd.value = f'=TEXT({START_CELL}+{idx},"ddd")'
    wd.font = FONT_LBL
    wd.fill = fill(BLUE_LIGHT)
    wd.alignment = Alignment(horizontal="center")
    wd.border = BORDER_ALL
    dn = hab.cell(row=6, column=col)
    dn.value = f'=DAY({START_CELL}+{idx})'
    dn.font = Font(name="Roboto", size=9, bold=True, color=GRAY_TXT)
    dn.fill = fill(CREAM)
    dn.alignment = Alignment(horizontal="center")
    dn.border = BORDER_ALL

# Habit rows 7-30 (24 rows: 8 examples + 16 blank)
HAB_FIRST_ROW, HAB_LAST_ROW = 7, 30
example_habits = [
    "💧 2 литра воды", "🏃 Спортзал в 19:00", "📖 Чтение 30 мин", "📓 Ведение дневника",
    "💼 Работа — 5 часов", "💪 Подработка — 2 часа", "🧘 Растяжка 15 мин", "🙏 Практика благодарности",
]
for i, r in enumerate(range(HAB_FIRST_ROW, HAB_LAST_ROW + 1)):
    hab.cell(row=r, column=1, value=i + 1).font = FONT_BODY
    name_cell = hab.cell(row=r, column=2)
    if i < len(example_habits):
        name_cell.value = example_habits[i]
    name_cell.font = FONT_BODY
    name_cell.border = BORDER_ALL
    day_range = f"{get_column_letter(FIRST_DAY_COL)}{r}:{get_column_letter(LAST_DAY_COL)}{r}"
    for col in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
        c = hab.cell(row=r, column=col, value=False)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER_ALL
    add_checkbox_validation(hab, day_range)
    hab.conditional_formatting.add(
        day_range,
        FormulaRule(formula=[f"{get_column_letter(FIRST_DAY_COL)}{r}=TRUE"], fill=fill(SAGE))
    )

NAME_RANGE = f"$B${HAB_FIRST_ROW}:$B${HAB_LAST_ROW}"

# Summary rows: Выполнены / Невыполнены / Прогресс %
ROW_DONE, ROW_NOTDONE, ROW_PCT = HAB_LAST_ROW + 1, HAB_LAST_ROW + 2, HAB_LAST_ROW + 3
hab.cell(row=ROW_DONE, column=2, value="ВЫПОЛНЕНЫ").font = FONT_LBL
hab.cell(row=ROW_NOTDONE, column=2, value="НЕВЫПОЛНЕНЫ").font = FONT_LBL
hab.cell(row=ROW_PCT, column=2, value="ПРОГРЕСС %").font = FONT_LBL
for r in (ROW_DONE, ROW_NOTDONE, ROW_PCT):
    hab.cell(row=r, column=2).fill = fill(GRAY_LIGHT)
    hab.cell(row=r, column=2).border = BORDER_ALL

for col in range(FIRST_DAY_COL, LAST_DAY_COL + 1):
    letter = get_column_letter(col)
    d = hab.cell(row=ROW_DONE, column=col)
    d.value = f'=COUNTIFS({NAME_RANGE},"<>",{letter}{HAB_FIRST_ROW}:{letter}{HAB_LAST_ROW},TRUE)'
    nd = hab.cell(row=ROW_NOTDONE, column=col)
    nd.value = f'=COUNTIFS({NAME_RANGE},"<>",{letter}{HAB_FIRST_ROW}:{letter}{HAB_LAST_ROW},FALSE)'
    p = hab.cell(row=ROW_PCT, column=col)
    p.value = f'=IFERROR({letter}{ROW_DONE}/({letter}{ROW_DONE}+{letter}{ROW_NOTDONE}),0)'
    p.number_format = "0%"
    for cell in (d, nd, p):
        cell.font = Font(name="Roboto", size=9, color=GRAY_TXT)
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER_ALL

pct_range = f"{get_column_letter(FIRST_DAY_COL)}{ROW_PCT}:{get_column_letter(LAST_DAY_COL)}{ROW_PCT}"
hab.conditional_formatting.add(
    pct_range,
    ColorScaleRule(start_type="min", start_color="FBEDF1", end_type="max", end_color=SAGE_DK)
)

hab.freeze_panes = f"C{HAB_FIRST_ROW}"

# Weekly habits table (10 rows)
WEEKLY_HDR_ROW = ROW_PCT + 2
box_header(hab, f"B{WEEKLY_HDR_ROW}:B{WEEKLY_HDR_ROW}", "ЕЖЕНЕДЕЛЬНЫЕ ПРИВЫЧКИ", BLUE)
for i, wk in enumerate(["НЕДЕЛЯ 1", "НЕДЕЛЯ 2", "НЕДЕЛЯ 3", "НЕДЕЛЯ 4", "НЕДЕЛЯ 5"]):
    box_header(hab, f"{get_column_letter(3+i)}{WEEKLY_HDR_ROW}:{get_column_letter(3+i)}{WEEKLY_HDR_ROW}", wk, BLUE_LIGHT)

weekly_examples = ["🏋 Спортзал 3 раза в неделю", "💰 Анализ расходов"]
WEEKLY_FIRST, WEEKLY_LAST = WEEKLY_HDR_ROW + 1, WEEKLY_HDR_ROW + 10
for i, r in enumerate(range(WEEKLY_FIRST, WEEKLY_LAST + 1)):
    nc = hab.cell(row=r, column=2)
    if i < len(weekly_examples):
        nc.value = weekly_examples[i]
    nc.font = FONT_BODY
    nc.border = BORDER_ALL
    for col in range(3, 8):
        cc = hab.cell(row=r, column=col, value=False)
        cc.alignment = Alignment(horizontal="center")
        cc.border = BORDER_ALL
add_checkbox_validation(hab, f"C{WEEKLY_FIRST}:G{WEEKLY_LAST}")
hab.conditional_formatting.add(
    f"C{WEEKLY_FIRST}:G{WEEKLY_LAST}",
    FormulaRule(formula=[f"C{WEEKLY_FIRST}=TRUE"], fill=fill(SAGE))
)

# Monthly habits table — placed BELOW the weekly table, not beside it (10 rows)
MONTHLY_HDR_ROW = WEEKLY_LAST + 2
box_header(hab, f"B{MONTHLY_HDR_ROW}:C{MONTHLY_HDR_ROW}", "ЕЖЕМЕСЯЧНЫЕ ПРИВЫЧКИ", PINK)
monthly_examples = ["📏 Сделать замеры до/после"]
MONTHLY_FIRST, MONTHLY_LAST = MONTHLY_HDR_ROW + 1, MONTHLY_HDR_ROW + 10
for i, r in enumerate(range(MONTHLY_FIRST, MONTHLY_LAST + 1)):
    nc = hab.cell(row=r, column=2)
    if i < len(monthly_examples):
        nc.value = monthly_examples[i]
    nc.font = FONT_BODY
    nc.border = BORDER_ALL
    cc = hab.cell(row=r, column=3, value=False)
    cc.alignment = Alignment(horizontal="center")
    cc.border = BORDER_ALL
add_checkbox_validation(hab, f"C{MONTHLY_FIRST}:C{MONTHLY_LAST}")
hab.conditional_formatting.add(
    f"C{MONTHLY_FIRST}:C{MONTHLY_LAST}",
    FormulaRule(formula=[f"C{MONTHLY_FIRST}=TRUE"], fill=fill(PINK))
)

# Notes section
NOTES_HDR_ROW = MONTHLY_LAST + 2
box_header(hab, f"B{NOTES_HDR_ROW}:P{NOTES_HDR_ROW}", "ЗАМЕТКИ", BLUE)
for r in range(NOTES_HDR_ROW + 1, NOTES_HDR_ROW + 7):
    for col in range(2, 17):
        c = hab.cell(row=r, column=col)
        c.border = Border(bottom=THIN)

# =====================================================================
# SHEET: ПЛАНИРОВЩИК ЗАДАЧ
# =====================================================================
tsk = wb.create_sheet("Планировщик задач")
tsk.sheet_properties.tabColor = BLUE_DK
tsk.sheet_view.showGridLines = False

set_col_widths(tsk, {"A": 3, "B": 32, "C": 12, "D": 8, "E": 13, "F": 7, "G": 12,
                      "H": 3, "I": 15, "J": 8, "K": 3, "L": 22, "M": 10, "N": 15, "O": 8})

style_title(tsk, "A1:G1", "🗂  МОИ ЗАДАЧИ", BLUE)
tsk.row_dimensions[1].height = 28

# ---- month/year filter (applies to the dashboard below, not the task list itself) ----
tsk["I2"] = "Месяц:"
tsk["I2"].font = FONT_LBL
tsk["J2"] = "Июль"
tsk["J2"].font = FONT_H2
tsk["J2"].fill = fill(YELLOW_LIGHT)
tsk["J2"].alignment = Alignment(horizontal="center")
tsk["J2"].border = BORDER_ALL
add_list_validation(tsk, "J2", RANGE_MONTHS)

tsk["K2"] = "Год:"
tsk["K2"].font = FONT_LBL
tsk["L2"] = 2026
tsk["L2"].font = FONT_H2
tsk["L2"].fill = fill(YELLOW_LIGHT)
tsk["L2"].alignment = Alignment(horizontal="center")
tsk["L2"].border = BORDER_ALL
add_list_validation(tsk, "L2", RANGE_YEARS)

TASK_MONTH_CELL, TASK_YEAR_CELL = "$J$2", "$L$2"
TASK_MONTH_NUM = f"MATCH({TASK_MONTH_CELL},{RANGE_MONTHS},0)"

TASK_HDR_ROW = 3
for col, label in zip("BCDEFG", ["Задача", "Срок", "Дни ⏳", "Приоритет", "✅", "Категория"]):
    c = tsk[f"{col}{TASK_HDR_ROW}"]
    c.value = label
    c.font = FONT_LBL
    c.fill = fill(BLUE_LIGHT)
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER_ALL

TASK_FIRST, TASK_LAST = TASK_HDR_ROW + 1, TASK_HDR_ROW + 20
task_names = [
    "🚶 10k шагов", "📕 20 страниц книги", "✉️ Отправить письмо клиенту", "🧘 Медитация (10 мин)",
    "💰 Посчитать бюджет на месяц", "🎬 Сделать рилс", "💳 Оплатить кредит", "🗓 План на следующую неделю",
    "📞 Позвонить маме", "🎤 Подготовить речь для конфы", "📊 Бюджет на месяц (план)", "📈 Составить отчёт",
]
task_offsets = [-3, -1, 0, 0, 2, 3, -2, 1, 0, 2, 10, -4]
task_priorities = ["Высокий", "Срочно", "Средний", "Высокий", "Низкий", "Срочно",
                    "Срочно", "Средний", "Средний", "Высокий", "Низкий", "Высокий"]
task_categories = ["Личное", "Личное", "Работа", "Личное", "Другое", "Работа",
                    "Другое", "Работа", "Личное", "Личное", "Личное", "Работа"]
task_done = [False, False, True, False, False, False, True, False, True, False, False, False]

PRIORITY_COLORS = {"Срочно": "F2B8C2", "Высокий": PURPLE, "Средний": BLUE, "Низкий": SAGE}

for i, r in enumerate(range(TASK_FIRST, TASK_LAST + 1)):
    if i < len(task_names):
        tsk.cell(row=r, column=2, value=task_names[i]).font = FONT_BODY
        dcell = tsk.cell(row=r, column=3, value=f"=TODAY()+({task_offsets[i]})")
        dcell.number_format = "dd.mm.yyyy"
        dcell.font = FONT_BODY
        dcell.alignment = Alignment(horizontal="center")
        ddays = tsk.cell(row=r, column=4, value=f"=C{r}-TODAY()")
        ddays.font = FONT_BODY
        ddays.alignment = Alignment(horizontal="center")
        pcell = tsk.cell(row=r, column=5, value=task_priorities[i])
        pcell.font = FONT_BODY
        pcell.alignment = Alignment(horizontal="center")
        fcell = tsk.cell(row=r, column=6, value=task_done[i])
        fcell.alignment = Alignment(horizontal="center")
        ccell = tsk.cell(row=r, column=7, value=task_categories[i])
        ccell.font = FONT_BODY
        ccell.alignment = Alignment(horizontal="center")
    else:
        tsk.cell(row=r, column=2).font = FONT_BODY
        tsk.cell(row=r, column=6, value=False)
    for col in range(2, 8):
        tsk.cell(row=r, column=col).border = BORDER_ALL

TASK_RANGE_ALL = f"B{TASK_FIRST}:G{TASK_LAST}"
add_list_validation(tsk, f"E{TASK_FIRST}:E{TASK_LAST}", RANGE_PRIORITY)
add_list_validation(tsk, f"G{TASK_FIRST}:G{TASK_LAST}", RANGE_TASKCAT)
add_checkbox_validation(tsk, f"F{TASK_FIRST}:F{TASK_LAST}")

for prio, color in PRIORITY_COLORS.items():
    tsk.conditional_formatting.add(
        f"E{TASK_FIRST}:E{TASK_LAST}",
        FormulaRule(formula=[f'E{TASK_FIRST}="{prio}"'], fill=fill(color))
    )
tsk.conditional_formatting.add(
    TASK_RANGE_ALL,
    FormulaRule(formula=[f"$F{TASK_FIRST}=TRUE"], font=Font(name="Roboto", size=10, strike=True, color="ABABAB"))
)
tsk.conditional_formatting.add(
    f"D{TASK_FIRST}:D{TASK_LAST}",
    CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Roboto", size=10, bold=True, color="C0392B"))
)

# ---- helpers: count tasks whose Срок falls in the selected Месяц/Год ----
def task_month_bare(extra_cond=None):
    date_cond = (f'(YEAR($C${TASK_FIRST}:$C${TASK_LAST}+0)={TASK_YEAR_CELL})'
                 f'*(MONTH($C${TASK_FIRST}:$C${TASK_LAST}+0)={TASK_MONTH_NUM})')
    if extra_cond:
        return f'SUMPRODUCT({extra_cond}*{date_cond})'
    return f'SUMPRODUCT({date_cond})'

def task_month_formula(extra_cond=None):
    return f'={task_month_bare(extra_cond)}'

# ---- summary stat rows (задачи со сроком в выбранном месяце) ----
STAT_ROW1 = TASK_LAST + 2
prio_labels = ["СРОЧНЫЕ", "ВЫСОКИЙ ПРИОРИТЕТ", "СРЕДНИЙ ПРИОРИТЕТ", "НИЗКИЙ ПРИОРИТЕТ"]
prio_colors_stat = ["F2B8C2", PURPLE, BLUE, SAGE]
for i, (lbl, color) in enumerate(zip(prio_labels, prio_colors_stat)):
    r = STAT_ROW1 + i
    lc = tsk.cell(row=r, column=2, value=lbl)
    lc.font = FONT_LBL
    lc.fill = fill(color)
    lc.border = BORDER_ALL
    prio_word = ["Срочно", "Высокий", "Средний", "Низкий"][i]
    vc = tsk.cell(row=r, column=3, value=task_month_formula(f'($E${TASK_FIRST}:$E${TASK_LAST}="{prio_word}")'))
    vc.font = FONT_H2
    vc.fill = fill(color)
    vc.alignment = Alignment(horizontal="center")
    vc.border = BORDER_ALL

cat_labels = ["ВСЕГО ЗАДАЧ", "ЛИЧНЫЕ", "РАБОТА", "ДРУГИЕ"]
for i, lbl in enumerate(cat_labels):
    r = STAT_ROW1 + i
    tsk.merge_cells(f"E{r}:F{r}")
    lc = tsk.cell(row=r, column=5, value=lbl)
    lc.font = FONT_LBL
    lc.fill = fill(BLUE_LIGHT)
    lc.alignment = Alignment(horizontal="center")
    vc = tsk.cell(row=r, column=7)
    if i == 0:
        vc.value = task_month_formula()
    else:
        word = {1: "Личное", 2: "Работа", 3: "Другое"}[i]
        vc.value = task_month_formula(f'($G${TASK_FIRST}:$G${TASK_LAST}="{word}")')
    vc.font = FONT_H2
    vc.fill = fill(BLUE_LIGHT)
    vc.alignment = Alignment(horizontal="center")
    for cc in (lc, vc):
        cc.border = BORDER_ALL

# filter control
FILTER_ROW = STAT_ROW1 + 5
tsk.cell(row=FILTER_ROW, column=2, value="Выделение задач:").font = FONT_LBL
filt_cell = tsk.cell(row=FILTER_ROW, column=3, value="(Все)")
filt_cell.fill = fill(YELLOW_LIGHT)
filt_cell.border = BORDER_ALL
filt_cell.alignment = Alignment(horizontal="center")
dv_filter = DataValidation(type="list", formula1='"(Все),Работа,Личное,Другое"', allow_blank=True)
tsk.add_data_validation(dv_filter)
dv_filter.add(f"C{FILTER_ROW}")
tsk.conditional_formatting.add(
    TASK_RANGE_ALL,
    FormulaRule(formula=[f'AND($C${FILTER_ROW}<>"(Все)",$G{TASK_FIRST}=$C${FILTER_ROW})'], fill=fill(YELLOW_LIGHT))
)

# stat tiles: done/progress (за выбранный месяц) + deadline-today/overdue (всегда "сегодня")
TILE_ROW = STAT_ROW1
DONE_MONTH_EXPR = task_month_bare(f'($F${TASK_FIRST}:$F${TASK_LAST}=TRUE)')
TOTAL_MONTH_EXPR = task_month_bare()
tsk.cell(row=TILE_ROW, column=12, value="ЗАДАЧ ВЫПОЛНЕНО (месяц)").font = FONT_LBL
tile_done = tsk.cell(row=TILE_ROW, column=13, value=f'={DONE_MONTH_EXPR}&"/"&{TOTAL_MONTH_EXPR}')
tsk.cell(row=TILE_ROW + 1, column=12, value="ПРОГРЕСС (месяц)").font = FONT_LBL
tile_pct = tsk.cell(row=TILE_ROW + 1, column=13, value=f'=IFERROR({DONE_MONTH_EXPR}/{TOTAL_MONTH_EXPR},0)')
tile_pct.number_format = "0%"
tsk.cell(row=TILE_ROW + 2, column=12, value="ДЕДЛАЙН СЕГОДНЯ").font = FONT_LBL
tsk.cell(row=TILE_ROW + 2, column=13, value=f'=COUNTIF($C${TASK_FIRST}:$C${TASK_LAST},TODAY())')
tsk.cell(row=TILE_ROW + 3, column=12, value="ПРОСРОЧЕННЫЕ ЗАДАЧИ").font = FONT_LBL
tile_over = tsk.cell(row=TILE_ROW + 3, column=13,
                      value=f'=COUNTIFS($C${TASK_FIRST}:$C${TASK_LAST},"<"&TODAY(),$F${TASK_FIRST}:$F${TASK_LAST},FALSE)')
for i in range(4):
    lbl_c = tsk.cell(row=TILE_ROW + i, column=12)
    val_c = tsk.cell(row=TILE_ROW + i, column=13)
    lbl_c.fill = fill(CREAM)
    val_c.fill = fill(CREAM)
    val_c.font = FONT_H2
    val_c.alignment = Alignment(horizontal="center")
    lbl_c.border = BORDER_ALL
    val_c.border = BORDER_ALL
tsk.conditional_formatting.add(
    f"M{TILE_ROW+3}",
    CellIsRule(operator="greaterThan", formula=["0"], fill=fill("F2B8C2"), font=Font(bold=True, color="7A2530"))
)

# priority helper table (for pie chart) — за выбранный месяц
PH_ROW = STAT_ROW1
for i, (prio, color) in enumerate(PRIORITY_COLORS.items()):
    r = PH_ROW + i
    tsk.cell(row=r, column=9, value=prio).font = FONT_BODY
    tsk.cell(row=r, column=10,
             value=task_month_formula(f'($E${TASK_FIRST}:$E${TASK_LAST}="{prio}")')).font = FONT_BODY

# category helper table (for bar chart) — за выбранный месяц
CH_ROW = STAT_ROW1
for i, word in enumerate(["Работа", "Личное", "Другое"]):
    r = CH_ROW + i
    tsk.cell(row=r, column=14, value=word).font = FONT_BODY
    tsk.cell(row=r, column=15,
             value=task_month_formula(f'($G${TASK_FIRST}:$G${TASK_LAST}="{word}")')).font = FONT_BODY

pie = PieChart()
pie.title = "Задачи по приоритету"
data = Reference(tsk, min_col=10, min_row=PH_ROW, max_row=PH_ROW + 3)
cats = Reference(tsk, min_col=9, min_row=PH_ROW, max_row=PH_ROW + 3)
pie.add_data(data, titles_from_data=False)
pie.set_categories(cats)
pie.height, pie.width = 7, 9
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
# charts anchored in a dedicated column, clear of every table on this sheet
CHART_COL = "R"
tsk.add_chart(pie, f"{CHART_COL}3")

bar = BarChart()
bar.type = "col"
bar.title = "Задачи по категориям"
bdata = Reference(tsk, min_col=15, min_row=CH_ROW, max_row=CH_ROW + 2)
bcats = Reference(tsk, min_col=14, min_row=CH_ROW, max_row=CH_ROW + 2)
bar.add_data(bdata, titles_from_data=False)
bar.set_categories(bcats)
bar.height, bar.width = 7, 9
bar.legend = None
tsk.add_chart(bar, f"{CHART_COL}20")

# ---- FILTER tables by priority (2x2 grid) ----
FT_ROW1 = FILTER_ROW + 17
filter_blocks = [
    ("СРОЧНЫЕ", "Срочно", "F2B8C2", 2),
    ("ВЫСОКИЙ ПРИОРИТЕТ", "Высокий", PURPLE, 7),
]
filter_blocks2 = [
    ("СРЕДНИЙ ПРИОРИТЕТ", "Средний", BLUE, 2),
    ("НИЗКИЙ ПРИОРИТЕТ", "Низкий", SAGE, 7),
]

def build_filter_block(ws, title, prio_word, color, start_col, header_row):
    c0 = get_column_letter(start_col)
    c3 = get_column_letter(start_col + 3)
    ws.merge_cells(f"{c0}{header_row}:{c3}{header_row}")
    hc = ws[f"{c0}{header_row}"]
    hc.value = title
    hc.font = FONT_LBL
    hc.fill = fill(color)
    hc.alignment = Alignment(horizontal="center")
    for row in ws[f"{c0}{header_row}:{c3}{header_row}"]:
        for c in row:
            c.fill = fill(color)
            c.border = BORDER_ALL
    sub_row = header_row + 1
    for i, lbl in enumerate(["Задача", "Срок", "Дни", "Тип"]):
        cc = ws.cell(row=sub_row, column=start_col + i, value=lbl)
        cc.font = FONT_LBL
        cc.fill = fill(BLUE_LIGHT if color != BLUE else BLUE_DK)
        cc.alignment = Alignment(horizontal="center")
        cc.border = BORDER_ALL
    data_row = sub_row + 1
    crit_rng = f"$E${TASK_FIRST}:$E${TASK_LAST}"
    date_cond = (f'(YEAR($C${TASK_FIRST}:$C${TASK_LAST}+0)={TASK_YEAR_CELL})'
                 f'*(MONTH($C${TASK_FIRST}:$C${TASK_LAST}+0)={TASK_MONTH_NUM})')
    match_cond = f'({crit_rng}="{prio_word}")*{date_cond}'
    formulas = [
        f'=IFERROR(FILTER($B${TASK_FIRST}:$B${TASK_LAST},{match_cond}),"")',
        f'=IFERROR(FILTER($C${TASK_FIRST}:$C${TASK_LAST},{match_cond}),"")',
        f'=IFERROR(FILTER($D${TASK_FIRST}:$D${TASK_LAST},{match_cond}),"")',
        f'=IFERROR(FILTER($G${TASK_FIRST}:$G${TASK_LAST},{match_cond}),"")',
    ]
    for i, f_ in enumerate(formulas):
        cell = ws.cell(row=data_row, column=start_col + i, value=f_)
        cell.font = FONT_BODY
    for r in range(data_row, data_row + 4):
        for col in range(start_col, start_col + 4):
            ws.cell(row=r, column=col).border = BORDER_ALL
    return data_row + 4

for title, prio_word, color, scol in filter_blocks:
    next_row = build_filter_block(tsk, title, prio_word, color, scol, FT_ROW1)

FT_ROW2 = FT_ROW1 + 8
for title, prio_word, color, scol in filter_blocks2:
    build_filter_block(tsk, title, prio_word, color, scol, FT_ROW2)

tsk.freeze_panes = f"B{TASK_FIRST}"

# =====================================================================
# SHEET: ЖУРНАЛ ОПЕРАЦИЙ (ввод данных)
# =====================================================================
jrn = wb.create_sheet("Журнал операций")
jrn.sheet_properties.tabColor = SAGE_DK
jrn.sheet_view.showGridLines = False

set_col_widths(jrn, {"A": 3, "B": 18, "C": 16, "D": 12, "E": 20, "F": 12, "G": 22, "H": 12})

style_title(jrn, "A1:H1", "📒  ЖУРНАЛ ОПЕРАЦИЙ", SAGE)
jrn.row_dimensions[1].height = 28

# ---- месяц/год выбираются на вкладке «Бюджет» (там главный переключатель) ----
jrn["B3"] = "Месяц:"
jrn["B3"].font = FONT_LBL
jrn["C3"] = "=Бюджет!$C$3"
jrn["C3"].font = FONT_H2
jrn["C3"].fill = fill(GRAY_LIGHT)
jrn["C3"].alignment = Alignment(horizontal="center")
jrn["C3"].border = BORDER_ALL

jrn["E3"] = "Год:"
jrn["E3"].font = FONT_LBL
jrn["F3"] = "=Бюджет!$F$3"
jrn["F3"].font = FONT_H2
jrn["F3"].fill = fill(GRAY_LIGHT)
jrn["F3"].alignment = Alignment(horizontal="center")
jrn["F3"].border = BORDER_ALL
jrn["G3"] = "← изменить можно на вкладке «Бюджет»"
jrn["G3"].font = FONT_NOTE

MONTH_CELL, YEAR_CELL = "Бюджет!$C$3", "Бюджет!$F$3"
MONTH_NUM = f"MATCH({MONTH_CELL},{RANGE_MONTHS},0)"

OPS_FIRST, OPS_LAST = 25, 224
B_RANGE = f"$B${OPS_FIRST}:$B${OPS_LAST}"   # Дата
C_RANGE = f"$C${OPS_FIRST}:$C${OPS_LAST}"   # Счёт
D_RANGE = f"$D${OPS_FIRST}:$D${OPS_LAST}"   # Тип
E_RANGE = f"$E${OPS_FIRST}:$E${OPS_LAST}"   # Категория
F_RANGE = f"$F${OPS_FIRST}:$F${OPS_LAST}"   # Сумма

def month_sum(txn, exclude_service=False):
    parts = [f'({D_RANGE}="{txn}")', f'(YEAR({B_RANGE}+0)={YEAR_CELL})', f'(MONTH({B_RANGE}+0)={MONTH_NUM})']
    if exclude_service:
        parts.append(f'(COUNTIF({RANGE_SERVICE},{E_RANGE})=0)')
    return f'=SUMPRODUCT({"*".join(parts)}*{F_RANGE})'

# ---- итоги: доход / расход / баланс за выбранный месяц ----
box_header(jrn, "B5:C5", "ДОХОД (месяц)", SAGE)
box_header(jrn, "D5:E5", "РАСХОД (месяц)", PINK)
box_header(jrn, "F5:H5", "БАЛАНС (месяц)", PURPLE)

jrn.merge_cells("B6:C8")
jrn["B6"] = month_sum("Доход", exclude_service=True)
jrn["B6"].number_format = '#,##0" ₽"'
jrn["B6"].font = FONT_BIG
jrn["B6"].alignment = Alignment(horizontal="center", vertical="center")

jrn.merge_cells("D6:E8")
jrn["D6"] = month_sum("Расход", exclude_service=True)
jrn["D6"].number_format = '#,##0" ₽"'
jrn["D6"].font = FONT_BIG
jrn["D6"].alignment = Alignment(horizontal="center", vertical="center")

jrn.merge_cells("F6:H8")
jrn["F6"] = "=B6-D6"
jrn["F6"].number_format = '#,##0" ₽"'
jrn["F6"].font = FONT_BIG
jrn["F6"].alignment = Alignment(horizontal="center", vertical="center")

for rng in ("B6:C8", "D6:E8", "F6:H8"):
    for row in jrn[rng]:
        for c in row:
            c.border = BORDER_ALL

# ---- остаток по счетам (не зависит от месяца — текущее состояние) ----
box_header(jrn, "B10:C10", "ОСТАТОК ПО СЧЕТАМ", BLUE)
jrn["B11"] = "Счёт"
jrn["C11"] = "Остаток"
for col in "BC":
    jrn[f"{col}11"].font = FONT_LBL
    jrn[f"{col}11"].fill = fill(BLUE_LIGHT)
    jrn[f"{col}11"].alignment = Alignment(horizontal="center")
    jrn[f"{col}11"].border = BORDER_ALL

ACC_FIRST = 12
N_ACC_SLOTS = 10
for i in range(N_ACC_SLOTS):
    r = ACC_FIRST + i
    acc_ref = f"Справочник!$K${3+i}"
    jrn[f"B{r}"] = f'=IFERROR(IF({acc_ref}="","",{acc_ref}),"")'
    jrn[f"C{r}"] = (f'=IF(B{r}="","",SUMIFS({F_RANGE},{C_RANGE},B{r},{D_RANGE},"Доход")'
                     f'-SUMIFS({F_RANGE},{C_RANGE},B{r},{D_RANGE},"Расход"))')
    jrn[f"B{r}"].font = FONT_BODY
    jrn[f"C{r}"].font = FONT_BODY
    jrn[f"C{r}"].number_format = '#,##0" ₽";[RED]-#,##0" ₽"'
    jrn[f"C{r}"].alignment = Alignment(horizontal="center")
    for col in "BC":
        jrn[f"{col}{r}"].border = BORDER_ALL
jrn.conditional_formatting.add(
    f"C{ACC_FIRST}:C{ACC_FIRST+N_ACC_SLOTS-1}",
    CellIsRule(operator="lessThan", formula=["0"], fill=fill(PINK_LIGHT))
)

# ---- журнал операций (основная таблица для ввода) ----
JOURNAL_HDR_ROW = ACC_FIRST + N_ACC_SLOTS + 1
box_header(jrn, f"B{JOURNAL_HDR_ROW}:H{JOURNAL_HDR_ROW}", "✏️  ВНОСИТЕ ДОХОДЫ И РАСХОДЫ СЮДА", SAGE)
JOURNAL_SUB_ROW = JOURNAL_HDR_ROW + 1
for col, lbl in zip("BCDEFGH", ["Дата", "Счёт", "Тип", "Категория", "Сумма", "Описание", "Остаток"]):
    c = jrn[f"{col}{JOURNAL_SUB_ROW}"]
    c.value = lbl
    c.font = FONT_LBL
    c.fill = fill(SAGE_LIGHT)
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER_ALL

example_ops = [
    (datetime.date(2026, 1, 5), "Карта Сбербанк", "Доход", "Зарплата", 75000, ""),
    (datetime.date(2026, 1, 6), "Карта Сбербанк", "Расход", "Аренда", 28000, ""),
    (datetime.date(2026, 1, 8), "Наличные", "Расход", "Продукты", 7500, ""),
    (datetime.date(2026, 1, 10), "Карта Сбербанк", "Расход", "Пополнение вклада", 15000, "перевод на вклад"),
    (datetime.date(2026, 1, 10), "Вклад", "Доход", "Пополнение вклада", 15000, ""),
    (datetime.date(2026, 1, 15), "Карта Тинькофф", "Доход", "Заказы", 20000, "у друга"),
    (datetime.date(2026, 1, 18), "Карта Сбербанк", "Расход", "Кредит", 9000, ""),
    (datetime.date(2026, 1, 20), "Наличные", "Доход", "Взял в долг", 10000, ""),
    (datetime.date(2026, 1, 25), "Карта Сбербанк", "Расход", "Обучение", 10000, ""),
    (datetime.date(2026, 2, 2), "Карта Сбербанк", "Доход", "Зарплата", 75000, ""),
    (datetime.date(2026, 2, 3), "Карта Сбербанк", "Расход", "Аренда", 28000, ""),
    (datetime.date(2026, 2, 5), "Наличные", "Расход", "Продукты", 8000, ""),
    (datetime.date(2026, 2, 8), "Карта Сбербанк", "Расход", "Пополнение вклада", 20000, ""),
    (datetime.date(2026, 2, 8), "Вклад", "Доход", "Пополнение вклада", 20000, ""),
    (datetime.date(2026, 2, 12), "Карта Тинькофф", "Доход", "Заказы", 30000, ""),
    (datetime.date(2026, 2, 15), "Наличные", "Расход", "Вернул долг", 5000, ""),
    (datetime.date(2026, 2, 18), "Карта Сбербанк", "Расход", "Ипотека", 18000, ""),
]

for i, r in enumerate(range(OPS_FIRST, OPS_LAST + 1)):
    if i < len(example_ops):
        date_v, acc_v, txn_v, cat_v, sum_v, desc_v = example_ops[i]
        jrn[f"B{r}"] = date_v
        jrn[f"B{r}"].number_format = "dd.mm.yyyy"
        jrn[f"C{r}"] = acc_v
        jrn[f"D{r}"] = txn_v
        jrn[f"E{r}"] = cat_v
        jrn[f"F{r}"] = sum_v
        jrn[f"G{r}"] = desc_v
    jrn[f"H{r}"] = (f'=IF(C{r}="","",SUMPRODUCT(($C${OPS_FIRST}:C{r}=C{r})'
                     f'*(2*($D${OPS_FIRST}:D{r}="Доход")-1)*$F${OPS_FIRST}:F{r}))')
    for col in "BCDEFGH":
        jrn[f"{col}{r}"].font = FONT_BODY
        jrn[f"{col}{r}"].border = BORDER_ALL
    jrn[f"F{r}"].number_format = '#,##0" ₽"'
    jrn[f"H{r}"].number_format = '#,##0" ₽";[RED]-#,##0" ₽"'

add_list_validation(jrn, f"C{OPS_FIRST}:C{OPS_LAST}", RANGE_ACCOUNTS)
add_list_validation_literal(jrn, f"D{OPS_FIRST}:D{OPS_LAST}", "Доход,Расход")
add_list_validation(jrn, f"E{OPS_FIRST}:E{OPS_LAST}", RANGE_ALLCATS)

txn_colors = {"Доход": SAGE, "Расход": PINK}
for txn, color in txn_colors.items():
    jrn.conditional_formatting.add(
        f"D{OPS_FIRST}:D{OPS_LAST}",
        FormulaRule(formula=[f'D{OPS_FIRST}="{txn}"'], fill=fill(color))
    )

jrn.freeze_panes = f"B{JOURNAL_SUB_ROW + 1}"

# =====================================================================
# SHEET: БЮДЖЕТ (дашборд: итоги + категории + динамика, без журнала)
# =====================================================================
bud = wb.create_sheet("Бюджет")
bud.sheet_properties.tabColor = SAGE_DK
bud.sheet_view.showGridLines = False

# три колоночных блока одинаковой ширины (B:D, F:H, J:L) — таблица + своя
# диаграмма прямо под ней, без «плавающих» диаграмм в отдельной зоне
set_col_widths(bud, {"A": 3, "B": 16, "C": 12, "D": 9, "E": 3,
                      "F": 16, "G": 12, "H": 9, "I": 3, "J": 11, "K": 12, "L": 12})

style_title(bud, "A1:L1", "💰  БЮДЖЕТ: ДОХОДЫ И РАСХОДЫ", SAGE)
bud.row_dimensions[1].height = 28

JRN = "'Журнал операций'"

# ---- главный переключатель месяца/года (на «Журнале» он лишь отражается) ----
bud["B3"] = "Месяц:"
bud["B3"].font = FONT_LBL
bud["C3"] = "Январь"
bud["C3"].font = FONT_H2
bud["C3"].fill = fill(YELLOW_LIGHT)
bud["C3"].alignment = Alignment(horizontal="center")
bud["C3"].border = BORDER_ALL
add_list_validation(bud, "C3", RANGE_MONTHS)

bud["D3"] = "Год:"
bud["D3"].font = FONT_LBL
bud["F3"] = 2026
bud["F3"].font = FONT_H2
bud["F3"].fill = fill(YELLOW_LIGHT)
bud["F3"].alignment = Alignment(horizontal="center")
bud["F3"].border = BORDER_ALL
add_list_validation(bud, "F3", RANGE_YEARS)

JB_RANGE = f"{JRN}!$B$25:$B$224"   # Дата
JC_RANGE = f"{JRN}!$C$25:$C$224"   # Счёт
JD_RANGE = f"{JRN}!$D$25:$D$224"   # Тип
JE_RANGE = f"{JRN}!$E$25:$E$224"   # Категория
JF_RANGE = f"{JRN}!$F$25:$F$224"   # Сумма
J_MONTH_CELL, J_YEAR_CELL = "$C$3", "$F$3"
J_MONTH_NUM = f"MATCH({J_MONTH_CELL},{RANGE_MONTHS},0)"

def bud_month_sum(txn, exclude_service=False):
    parts = [f'({JD_RANGE}="{txn}")', f'(YEAR({JB_RANGE}+0)={J_YEAR_CELL})', f'(MONTH({JB_RANGE}+0)={J_MONTH_NUM})']
    if exclude_service:
        parts.append(f'(COUNTIF({RANGE_SERVICE},{JE_RANGE})=0)')
    return f'=SUMPRODUCT({"*".join(parts)}*{JF_RANGE})'

# ---- KPI row, column-aligned with the 3 blocks below ----
box_header(bud, "B5:D5", "ДОХОД (месяц)", SAGE)
box_header(bud, "F5:H5", "РАСХОД (месяц)", PINK)
box_header(bud, "J5:L5", "БАЛАНС (месяц)", PURPLE)

bud.merge_cells("B6:D8")
bud["B6"] = bud_month_sum("Доход", exclude_service=True)
bud["B6"].number_format = '#,##0" ₽"'
bud["B6"].font = FONT_BIG
bud["B6"].alignment = Alignment(horizontal="center", vertical="center")

bud.merge_cells("F6:H8")
bud["F6"] = bud_month_sum("Расход", exclude_service=True)
bud["F6"].number_format = '#,##0" ₽"'
bud["F6"].font = FONT_BIG
bud["F6"].alignment = Alignment(horizontal="center", vertical="center")

bud.merge_cells("J6:L8")
bud["J6"] = "=B6-F6"
bud["J6"].number_format = '#,##0" ₽"'
bud["J6"].font = FONT_BIG
bud["J6"].alignment = Alignment(horizontal="center", vertical="center")

for rng in ("B6:D8", "F6:H8", "J6:L8"):
    for row in bud[rng]:
        for c in row:
            c.border = BORDER_ALL

# ---- три блока в ряд: Доходы по категориям | Расходы по категориям | Динамика за год ----
CAT_HDR_ROW = 10
box_header(bud, f"B{CAT_HDR_ROW}:D{CAT_HDR_ROW}", "ДОХОДЫ ПО КАТЕГОРИЯМ (месяц)", SAGE)
box_header(bud, f"F{CAT_HDR_ROW}:H{CAT_HDR_ROW}", "РАСХОДЫ ПО КАТЕГОРИЯМ (месяц)", PINK)
box_header(bud, f"J{CAT_HDR_ROW}:L{CAT_HDR_ROW}", "ДИНАМИКА ЗА ГОД", BLUE)

SUB_ROW = CAT_HDR_ROW + 1
for col, lbl in zip("BCD", ["Категория", "Сумма", "%"]):
    c = bud[f"{col}{SUB_ROW}"]
    c.value = lbl
    c.font = FONT_LBL
    c.fill = fill(SAGE_LIGHT)
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER_ALL
for col, lbl in zip("FGH", ["Категория", "Сумма", "%"]):
    c = bud[f"{col}{SUB_ROW}"]
    c.value = lbl
    c.font = FONT_LBL
    c.fill = fill(PINK_LIGHT)
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER_ALL
for col, lbl in zip("JKL", ["Месяц", "Доход", "Расход"]):
    c = bud[f"{col}{SUB_ROW}"]
    c.value = lbl
    c.font = FONT_LBL
    c.fill = fill(BLUE_LIGHT)
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER_ALL

DATA_ROW = SUB_ROW + 1
for i, cat in enumerate(income_cats):
    r = DATA_ROW + i
    bud[f"B{r}"] = cat
    bud[f"C{r}"] = (f'=SUMPRODUCT(({JD_RANGE}="Доход")*({JE_RANGE}="{cat}")'
                     f'*(YEAR({JB_RANGE}+0)={J_YEAR_CELL})*(MONTH({JB_RANGE}+0)={J_MONTH_NUM})*{JF_RANGE})')
    bud[f"D{r}"] = f'=IFERROR(C{r}/$B$6,0)'
    bud[f"D{r}"].number_format = "0%"
    for col in "BCD":
        bud[f"{col}{r}"].font = FONT_BODY
        bud[f"{col}{r}"].border = BORDER_ALL

for i, cat in enumerate(expense_cats):
    r = DATA_ROW + i
    bud[f"F{r}"] = cat
    bud[f"G{r}"] = (f'=SUMPRODUCT(({JD_RANGE}="Расход")*({JE_RANGE}="{cat}")'
                     f'*(YEAR({JB_RANGE}+0)={J_YEAR_CELL})*(MONTH({JB_RANGE}+0)={J_MONTH_NUM})*{JF_RANGE})')
    bud[f"H{r}"] = f'=IFERROR(G{r}/$F$6,0)'
    bud[f"H{r}"].number_format = "0%"
    for col in "FGH":
        bud[f"{col}{r}"].font = FONT_BODY
        bud[f"{col}{r}"].border = BORDER_ALL

for m in range(1, 13):
    r = DATA_ROW + m - 1
    bud[f"J{r}"] = f'=TEXT(DATE({J_YEAR_CELL},{m},1),"mmm yy")'
    bud[f"K{r}"] = (f'=SUMPRODUCT(({JD_RANGE}="Доход")*(YEAR({JB_RANGE}+0)={J_YEAR_CELL})'
                     f'*(MONTH({JB_RANGE}+0)={m})*(COUNTIF({RANGE_SERVICE},{JE_RANGE})=0)*{JF_RANGE})')
    bud[f"L{r}"] = (f'=SUMPRODUCT(({JD_RANGE}="Расход")*(YEAR({JB_RANGE}+0)={J_YEAR_CELL})'
                     f'*(MONTH({JB_RANGE}+0)={m})*(COUNTIF({RANGE_SERVICE},{JE_RANGE})=0)*{JF_RANGE})')
    for col in "JKL":
        bud[f"{col}{r}"].font = FONT_BODY
        bud[f"{col}{r}"].border = BORDER_ALL

INCOME_LAST = DATA_ROW + len(income_cats) - 1
EXPENSE_LAST = DATA_ROW + len(expense_cats) - 1
DYN_LAST = DATA_ROW + 11

# ---- диаграммы прямо под своей таблицей (та же колоночная зона) ----
CHART_ROW = max(EXPENSE_LAST, DYN_LAST) + 2

pie_income = PieChart()
pie_income.title = "Доходы по категориям"
pie_income.add_data(Reference(bud, min_col=3, min_row=DATA_ROW, max_row=INCOME_LAST))
pie_income.set_categories(Reference(bud, min_col=2, min_row=DATA_ROW, max_row=INCOME_LAST))
pie_income.height, pie_income.width = 7, 9
pie_income.legend.position = "b"
bud.add_chart(pie_income, f"B{CHART_ROW}")

pie_expense = PieChart()
pie_expense.title = "Расходы по категориям"
pie_expense.add_data(Reference(bud, min_col=7, min_row=DATA_ROW, max_row=EXPENSE_LAST))
pie_expense.set_categories(Reference(bud, min_col=6, min_row=DATA_ROW, max_row=EXPENSE_LAST))
pie_expense.height, pie_expense.width = 7, 9
pie_expense.legend.position = "b"
bud.add_chart(pie_expense, f"F{CHART_ROW}")

line = LineChart()
line.title = "Доход / расход по месяцам"
line.add_data(Reference(bud, min_col=11, min_row=SUB_ROW, max_row=DYN_LAST), titles_from_data=True)
line.add_data(Reference(bud, min_col=12, min_row=SUB_ROW, max_row=DYN_LAST), titles_from_data=True)
line.set_categories(Reference(bud, min_col=10, min_row=DATA_ROW, max_row=DYN_LAST))
line.height, line.width = 7, 9
bud.add_chart(line, f"J{CHART_ROW}")

# =====================================================================
# SHEET: КРЕДИТЫ (кредиты, кредитные карты и план их закрытия)
# =====================================================================
cr = wb.create_sheet("Кредиты")
cr.sheet_properties.tabColor = PURPLE_DK
cr.sheet_view.showGridLines = False

set_col_widths(cr, {"A": 3, "B": 18, "C": 14, "D": 14, "E": 12, "F": 12, "G": 8,
                     "H": 11, "I": 13, "J": 10, "K": 9, "L": 9, "M": 10, "N": 20})

style_title(cr, "A1:N1", "💳  КРЕДИТЫ И КРЕДИТНЫЕ КАРТЫ", PURPLE)
cr.row_dimensions[1].height = 28

# ---- стратегия закрытия ----
cr["B3"] = "Стратегия закрытия:"
cr["B3"].font = FONT_LBL
cr["C3"] = "Снежный ком"
cr["C3"].font = FONT_H2
cr["C3"].fill = fill(YELLOW_LIGHT)
cr["C3"].alignment = Alignment(horizontal="center")
cr["C3"].border = BORDER_ALL
add_list_validation(cr, "C3", RANGE_PAYOFF_STRATEGIES)
cr.merge_cells("E3:N3")
cr["E3"] = ("ℹ️ Снежный ком — сначала гасите кредит с наименьшим остатком (быстрые победы, мотивация). "
            "Лавина — сначала гасите кредит с наибольшей ставкой (математически выгоднее).")
cr["E3"].font = FONT_NOTE
cr["E3"].fill = fill(YELLOW_LIGHT)
cr["E3"].alignment = Alignment(wrap_text=True, vertical="center")

CR_FIRST, CR_LAST = 11, 20
F_RNG = f"$F${CR_FIRST}:$F${CR_LAST}"   # Остаток долга
G_RNG = f"$G${CR_FIRST}:$G${CR_LAST}"   # Ставка %
J_RNG = f"$J${CR_FIRST}:$J${CR_LAST}"   # Статус
M_RNG = f"$M${CR_FIRST}:$M${CR_LAST}"   # Осталось мес.

# ---- итоги ----
box_header(cr, "B5:C5", "ВСЕГО ОСТАТОК ДОЛГА", SAGE)
box_header(cr, "D5:E5", "ЕЖЕМЕСЯЧНЫЙ ПЛАТЁЖ", PINK)
box_header(cr, "F5:H5", "ПОЛНОЕ ПОГАШЕНИЕ ПРИМЕРНО", PURPLE)

cr.merge_cells("B6:C7")
cr["B6"] = f'=SUMIF($J${CR_FIRST}:$J${CR_LAST},"Активен",{F_RNG})'
cr["B6"].number_format = '#,##0" ₽"'
cr["B6"].font = FONT_BIG
cr["B6"].alignment = Alignment(horizontal="center", vertical="center")

cr.merge_cells("D6:E7")
cr["D6"] = f'=SUMIF($J${CR_FIRST}:$J${CR_LAST},"Активен",$H${CR_FIRST}:$H${CR_LAST})'
cr["D6"].number_format = '#,##0" ₽"'
cr["D6"].font = FONT_BIG
cr["D6"].alignment = Alignment(horizontal="center", vertical="center")

cr.merge_cells("F6:H7")
cr["F6"] = f'="≈ "&TEXT(EDATE(TODAY(),MAXIFS({M_RNG},$J${CR_FIRST}:$J${CR_LAST},"Активен")),"mmmm yyyy")'
cr["F6"].font = FONT_H2
cr["F6"].alignment = Alignment(horizontal="center", vertical="center")

for rng in ("B6:C7", "D6:E7", "F6:H7"):
    for row in cr[rng]:
        for c in row:
            c.border = BORDER_ALL

# ---- основная таблица кредитов ----
box_header(cr, "B9:N9", "КРЕДИТЫ И КРЕДИТНЫЕ КАРТЫ", PURPLE)
headers_cr = ["Счёт/Карта", "Тип", "Банк", "Сумма кредита", "Остаток долга", "Ставка %",
              "Платёж/мес", "Плановое закрытие", "Статус", "Приоритет", "Прогресс", "Осталось мес.", "Заметки"]
for col, lbl in zip("BCDEFGHIJKLMN", headers_cr):
    c = cr[f"{col}10"]
    c.value = lbl
    c.font = FONT_LBL
    c.fill = fill(PURPLE)
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    c.border = BORDER_ALL

example_credits = [
    ("Кредит 1", "Кредит", "Сбербанк", 300000, 180000, 12, 15000, datetime.date(2027, 6, 1), "Активен"),
    ("Кредит 2", "Кредит", "Тинькофф", 150000, 45000, 15, 8000, datetime.date(2026, 12, 1), "Активен"),
    ("Карта Тинькофф", "Кредитная карта", "Тинькофф", 100000, 32000, 24, 3000, datetime.date(2028, 1, 1), "Активен"),
    ("Кредит на телефон", "Кредит", "Сбербанк", 40000, 0, 10, 4000, datetime.date(2026, 3, 1), "Закрыт"),
]

for i, r in enumerate(range(CR_FIRST, CR_LAST + 1)):
    if i < len(example_credits):
        name_v, type_v, bank_v, sum_v, bal_v, rate_v, pay_v, date_v, status_v = example_credits[i]
        cr[f"B{r}"] = name_v
        cr[f"C{r}"] = type_v
        cr[f"D{r}"] = bank_v
        cr[f"E{r}"] = sum_v
        cr[f"F{r}"] = bal_v
        cr[f"G{r}"] = rate_v
        cr[f"H{r}"] = pay_v
        cr[f"I{r}"] = date_v
        cr[f"I{r}"].number_format = "dd.mm.yyyy"
        cr[f"J{r}"] = status_v
    cr[f"K{r}"] = (f'=IF($J{r}<>"Активен","",IF($C$3="Снежный ком",'
                    f'SUMPRODUCT(({F_RNG}<F{r})*($J${CR_FIRST}:$J${CR_LAST}="Активен"))+1,'
                    f'SUMPRODUCT(({G_RNG}>G{r})*($J${CR_FIRST}:$J${CR_LAST}="Активен"))+1))')
    cr[f"L{r}"] = f'=IFERROR((E{r}-F{r})/E{r},0)'
    cr[f"L{r}"].number_format = "0%"
    cr[f"M{r}"] = f'=IF(OR(F{r}="",H{r}="",H{r}=0),"",ROUNDUP(F{r}/H{r},0))'
    for col in "BCDEFGHIJKLMN":
        cell = cr[f"{col}{r}"]
        cell.font = FONT_BODY
        cell.border = BORDER_ALL
        cell.alignment = Alignment(horizontal="center") if col not in "BDN" else Alignment(horizontal="left")
    cr[f"E{r}"].number_format = '#,##0" ₽"'
    cr[f"F{r}"].number_format = '#,##0" ₽"'
    cr[f"H{r}"].number_format = '#,##0" ₽"'

add_list_validation(cr, f"C{CR_FIRST}:C{CR_LAST}", RANGE_CREDIT_TYPES)
add_list_validation(cr, f"B{CR_FIRST}:B{CR_LAST}", RANGE_ACCOUNTS)
add_list_validation(cr, f"J{CR_FIRST}:J{CR_LAST}", RANGE_CREDIT_STATUSES)

status_colors = {"Активен": SAGE, "Закрыт": GRAY_LIGHT}
for status, color in status_colors.items():
    cr.conditional_formatting.add(
        f"J{CR_FIRST}:J{CR_LAST}",
        FormulaRule(formula=[f'J{CR_FIRST}="{status}"'], fill=fill(color))
    )
cr.conditional_formatting.add(
    f"L{CR_FIRST}:L{CR_LAST}",
    DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=SAGE_DK)
)
cr.freeze_panes = f"B{CR_FIRST}"

# ---- диаграммы под таблицей: остаток по кредитам ----
CR_CHART_ROW = CR_LAST + 2
bar_cr = BarChart()
bar_cr.type = "col"
bar_cr.title = "Остаток долга по кредитам"
bar_cr.add_data(Reference(cr, min_col=6, min_row=CR_FIRST, max_row=CR_LAST))
bar_cr.set_categories(Reference(cr, min_col=2, min_row=CR_FIRST, max_row=CR_LAST))
bar_cr.height, bar_cr.width = 7, 10
bar_cr.legend = None
cr.add_chart(bar_cr, f"B{CR_CHART_ROW}")

pie_cr = PieChart()
pie_cr.title = "Доля в общем долге"
pie_cr.add_data(Reference(cr, min_col=6, min_row=CR_FIRST, max_row=CR_LAST))
pie_cr.set_categories(Reference(cr, min_col=2, min_row=CR_FIRST, max_row=CR_LAST))
pie_cr.height, pie_cr.width = 7, 10
cr.add_chart(pie_cr, f"J{CR_CHART_ROW}")

# ---- прогноз остатка по месяцам (при текущих платеже и ставке) ----
PROJ_HDR_ROW = CR_CHART_ROW + 15
box_header(cr, f"B{PROJ_HDR_ROW}:N{PROJ_HDR_ROW}", "ПРОГНОЗ ОСТАТКА ПО МЕСЯЦАМ (при текущих платежах)", BLUE)
PROJ_SUB_ROW = PROJ_HDR_ROW + 1
cr[f"B{PROJ_SUB_ROW}"] = "Счёт/Карта"
cr[f"B{PROJ_SUB_ROW}"].font = FONT_LBL
cr[f"B{PROJ_SUB_ROW}"].fill = fill(BLUE_LIGHT)
cr[f"B{PROJ_SUB_ROW}"].border = BORDER_ALL
for m in range(1, 13):
    col = 2 + m  # C..N
    letter = get_column_letter(col)
    c = cr[f"{letter}{PROJ_SUB_ROW}"]
    c.value = f'=TEXT(EDATE(TODAY(),{m}),"mmm yy")'
    c.font = FONT_LBL
    c.fill = fill(BLUE_LIGHT)
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER_ALL

PROJ_FIRST = PROJ_SUB_ROW + 1
PROJ_LAST = PROJ_FIRST + (CR_LAST - CR_FIRST)
for i in range(CR_LAST - CR_FIRST + 1):
    proj_row = PROJ_FIRST + i
    cr_row = CR_FIRST + i
    cr[f"B{proj_row}"] = f"=B{cr_row}"
    cr[f"B{proj_row}"].font = FONT_BODY
    cr[f"B{proj_row}"].border = BORDER_ALL
    for m in range(1, 13):
        col = 2 + m
        letter = get_column_letter(col)
        if m == 1:
            prev_ref = f"$F${cr_row}"
        else:
            prev_letter = get_column_letter(col - 1)
            prev_ref = f"{prev_letter}{proj_row}"
        cell = cr[f"{letter}{proj_row}"]
        cell.value = (f'=IF({prev_ref}="","",MAX(0,{prev_ref}*(1+$G${cr_row}/100/12)-$H${cr_row}))')
        cell.font = FONT_BODY
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER_ALL
        cell.number_format = '#,##0'

TOTAL_ROW = PROJ_LAST + 1
cr[f"B{TOTAL_ROW}"] = "ИТОГО"
cr[f"B{TOTAL_ROW}"].font = FONT_LBL
cr[f"B{TOTAL_ROW}"].fill = fill(GRAY_LIGHT)
cr[f"B{TOTAL_ROW}"].border = BORDER_ALL
for m in range(1, 13):
    col = 2 + m
    letter = get_column_letter(col)
    c = cr[f"{letter}{TOTAL_ROW}"]
    c.value = f'=SUM({letter}{PROJ_FIRST}:{letter}{PROJ_LAST})'
    c.font = FONT_LBL
    c.fill = fill(GRAY_LIGHT)
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER_ALL
    c.number_format = '#,##0'

# ---- диаграмма: траектория общего долга ----
PROJ_CHART_ROW = TOTAL_ROW + 2
line_cr = LineChart()
line_cr.title = "Прогноз общего остатка долга"
total_row_ref = Reference(cr, min_col=2, max_col=14, min_row=TOTAL_ROW, max_row=TOTAL_ROW)
line_cr.add_data(total_row_ref, titles_from_data=True, from_rows=True)
line_cr.set_categories(Reference(cr, min_col=3, max_col=14, min_row=PROJ_SUB_ROW, max_row=PROJ_SUB_ROW))
line_cr.height, line_cr.width = 8, 16
cr.add_chart(line_cr, f"B{PROJ_CHART_ROW}")

# =====================================================================
# final touches
# =====================================================================
wb._sheets = [bud, jrn, cr, hab, tsk, ref]
wb.active = 0

OUT_PATH = "Планировщик_Бюджет_Привычки.xlsx"
wb.save(OUT_PATH)
print(f"Saved {OUT_PATH}")
